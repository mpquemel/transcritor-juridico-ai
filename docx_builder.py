# Criado por Melhym Quemel <mpquemel@gmail.com>

import os
from docx import Document

def save_to_docx(formatted_text, output_dir, filename="Transcricao_Consolidada.docx"):
    doc = Document()
    doc.add_heading('Transcrição de Depoimentos - Formato MPRJ', 0)
    
    doc.add_paragraph(formatted_text)
    
    # Adicionando o aviso de responsabilidade obrigatório
    doc.add_heading('Aviso de Verificação Humana', level=1)
    doc.add_paragraph(
        "ATENÇÃO: A fidedignidade de cada palavra na transcrição final é de inteira responsabilidade "
        "do profissional. Utilize os arquivos .srt gerados na pasta de saída para rastreabilidade e "
        "conferência no áudio original."
    )
    
    filepath = os.path.join(output_dir, filename)
    doc.save(filepath)
    return filepath